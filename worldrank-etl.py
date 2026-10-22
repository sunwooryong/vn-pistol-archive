# -*- coding: utf-8 -*-
# ISSF 세계랭킹 자동 수집 → web/training-check/worldrank.json
#   ISSF가 매월 1일 공개하는 월간 랭킹 XLSX(World-Ranking-YYYYMMDD.xlsx)를 조회·파싱.
#   사용: python worldrank-etl.py
import os, re, json, io, datetime, subprocess, tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, 'web', 'training-check', 'worldrank.json')
PAGE = 'https://www.issf-sports.org/competitions/worldranking'
UA = {
    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36',
    'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8',
    'Accept-Language': 'en-US,en;q=0.9',
}
TOPN = 8

LABELS = {
    'APM':  {'ko': '10m 공기권총 남자', 'vi': '10m Súng ngắn hơi Nam', 'en': '10m Air Pistol Men'},
    'APW':  {'ko': '10m 공기권총 여자', 'vi': '10m Súng ngắn hơi Nữ', 'en': '10m Air Pistol Women'},
    'RFPM': {'ko': '25m 속사권총 남자', 'vi': '25m Bắn nhanh Nam', 'en': '25m Rapid Fire Pistol Men'},
    'SPW':  {'ko': '25m 권총 여자', 'vi': '25m Súng ngắn Nữ', 'en': '25m Pistol Women'},
    'ARM':  {'ko': '10m 공기소총 남자', 'vi': '10m Súng trường hơi Nam', 'en': '10m Air Rifle Men'},
    'ARW':  {'ko': '10m 공기소총 여자', 'vi': '10m Súng trường hơi Nữ', 'en': '10m Air Rifle Women'},
    'R3PM': {'ko': '50m 소총 3자세 남자', 'vi': '50m STr 3 tư thế Nam', 'en': '50m Rifle 3P Men'},
    'R3PW': {'ko': '50m 소총 3자세 여자', 'vi': '50m STr 3 tư thế Nữ', 'en': '50m Rifle 3P Women'},
    'TRM':  {'ko': '트랩 남자', 'vi': 'Trap Nam', 'en': 'Trap Men'},
    'TRW':  {'ko': '트랩 여자', 'vi': 'Trap Nữ', 'en': 'Trap Women'},
    'SKM':  {'ko': '스키트 남자', 'vi': 'Skeet Nam', 'en': 'Skeet Men'},
    'SKW':  {'ko': '스키트 여자', 'vi': 'Skeet Nữ', 'en': 'Skeet Women'},
}
ORDER = ['APM', 'APW', 'RFPM', 'SPW', 'ARM', 'ARW', 'R3PM', 'R3PW', 'TRM', 'TRW', 'SKM', 'SKW']


def fetch(url, binary=False):
    # urllib은 ISSF CDN에서 차단(404)되므로 curl 사용(Win10+/Git-Bash 기본 제공)
    if binary:
        tf = tempfile.NamedTemporaryFile(delete=False, suffix='.bin'); tf.close()
        subprocess.run(['curl', '-sL', '--fail', '-A', UA['User-Agent'], '-o', tf.name, url], check=True, timeout=90)
        data = open(tf.name, 'rb').read()
        os.unlink(tf.name)
        return data
    r = subprocess.run(['curl', '-sL', '--fail', '-A', UA['User-Agent'], url], check=True, capture_output=True, timeout=90)
    return r.stdout.decode('utf-8', 'replace')


def latest_xlsx_url():
    html = fetch(PAGE)
    links = re.findall(r'(https?://[^"\'<> ]*getfile\.aspx[^"\'<> ]*World-Ranking-(\d{8})\.xlsx)', html, re.I)
    if not links:
        rel = re.findall(r'([^"\'<> ]*getfile\.aspx[^"\'<> ]*World-Ranking-(\d{8})\.xlsx)', html, re.I)
        links = [(u if u.startswith('http') else 'https://backoffice.issf-sports.org/' + u.lstrip('/'), d) for u, d in rel]
    if not links:
        raise RuntimeError('랭킹 XLSX 링크를 찾지 못함')
    links.sort(key=lambda x: x[1])
    url = links[-1][0].replace('\\u0026', '&').replace('&amp;', '&')  # JS/HTML 이스케이프 복원
    return url, links[-1][1]


def build():
    import openpyxl
    url, datestr = latest_xlsx_url()
    buf = fetch(url, binary=True)
    wb = openpyxl.load_workbook(io.BytesIO(buf), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    idx = {str(v).strip().upper(): i for i, v in enumerate(header) if v is not None}
    iE, iR, iRt, iN, iNOC = idx.get('EVENT'), idx.get('RANK'), idx.get('RATING'), idx.get('NAME'), idx.get('NOC')
    if None in (iE, iR, iN):
        raise RuntimeError('헤더(EVENT/RANK/NAME) 누락: ' + repr(list(idx.keys())))

    events = {}
    for row in rows:
        ev = (str(row[iE]).strip() if row[iE] is not None else '')
        if not ev:
            continue
        try:
            rank = int(row[iR])
        except (TypeError, ValueError):
            continue
        if rank < 1 or rank > TOPN:
            continue
        name = (str(row[iN]).strip() if row[iN] is not None else '')
        noc = (str(row[iNOC]).strip() if (iNOC is not None and row[iNOC] is not None) else '')
        pts = row[iRt] if iRt is not None else None
        try:
            pts = int(pts)
        except (TypeError, ValueError):
            pass
        events.setdefault(ev, []).append({'rank': rank, 'name': name, 'noc': noc, 'points': pts})

    codes = [c for c in ORDER if c in events] + [c for c in events if c not in ORDER]
    out_events = [{
        'code': code,
        'label': LABELS.get(code, {'ko': code, 'vi': code, 'en': code}),
        'ranks': sorted(events[code], key=lambda x: x['rank'])[:TOPN],
    } for code in codes]

    ver = f'{datestr[:4]}.{datestr[4:6]}.{datestr[6:8]}'
    out = {
        'source': 'ISSF World Ranking',
        'version': ver,
        'generated_at': datetime.datetime.now().isoformat(timespec='seconds'),
        'events': out_events,
    }
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    json.dump(out, open(OUT, 'w', encoding='utf-8'), ensure_ascii=False)
    print(f'[worldrank.json] {len(out_events)} events, version {ver}')
    print('  ' + url)


if __name__ == '__main__':
    build()
